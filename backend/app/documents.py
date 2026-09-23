import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import load_workbook

CLAUSE_RE = re.compile(r"^(\d+(?:\.\d+){1,4})[.)]?\s*(.*)$")


@dataclass(frozen=True)
class Clause:
    document: str
    clause: str
    text: str
    page: int | None = None

    def render(self) -> str:
        page = f", стр. {self.page}" if self.page else ""
        return f"[{self.document}; п. {self.clause}{page}] {self.text}"


def display_name(path: Path) -> str:
    name = path.name
    return name[33:] if len(name) > 33 and name[32] == "_" else name


def _group_numbered(lines: list[tuple[str, int | None]], document: str) -> list[Clause]:
    clauses: list[Clause] = []
    current_number: str | None = None
    current_page: int | None = None
    current_parts: list[str] = []

    def flush() -> None:
        if current_number and current_parts:
            clauses.append(
                Clause(
                    document=document,
                    clause=current_number,
                    text=" ".join(current_parts).strip(),
                    page=current_page,
                )
            )

    for raw, page in lines:
        text = " ".join(raw.split())
        if not text:
            continue
        match = CLAUSE_RE.match(text)
        if match:
            flush()
            current_number = match.group(1)
            current_page = page
            current_parts = [match.group(2).strip()]
        elif current_number:
            current_parts.append(text)
    flush()
    return clauses


def parse_docx(path: Path) -> list[Clause]:
    document = Document(path)
    lines = [(paragraph.text, None) for paragraph in document.paragraphs]
    for table in document.tables:
        for row_index, row in enumerate(table.rows, 1):
            value = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if value:
                lines.append((f"Таблица, строка {row_index}: {value}", None))
    return _group_numbered(lines, display_name(path))


def parse_pdf(path: Path) -> list[Clause]:
    lines: list[tuple[str, int | None]] = []
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, 1):
            lines.extend((line, page_index) for line in page.get_text().splitlines())
    return _group_numbered(lines, display_name(path))


def parse_xlsx(path: Path) -> list[Clause]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    clauses: list[Clause] = []
    for sheet in workbook.worksheets:
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                clauses.append(
                    Clause(
                        document=display_name(path),
                        clause=f"{sheet.title}!{row_index}",
                        text=" | ".join(values),
                    )
                )
    return clauses


def parse_documents(paths: list[Path]) -> list[Clause]:
    parsers = {".docx": parse_docx, ".pdf": parse_pdf, ".xlsx": parse_xlsx}
    clauses: list[Clause] = []
    for path in paths:
        clauses.extend(parsers[path.suffix.lower()](path))
    return clauses
