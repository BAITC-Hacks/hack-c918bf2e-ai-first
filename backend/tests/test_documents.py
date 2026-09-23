from pathlib import Path

import pytest
from docx import Document

from app.documents import display_name, parse_documents, parse_docx, split_inline_clauses


def save_document(tmp_path: Path, document) -> Path:
    path = tmp_path / "regulation.docx"
    document.save(path)
    return path


def test_unnumbered_and_automatic_list_text_is_preserved(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("Департамент контроля")
    document.add_paragraph("Проверяет исполнение решений", style="List Number")
    document.add_paragraph("Ведёт реестр нарушений", style="List Number")
    clauses = parse_docx(save_document(tmp_path, document))
    assert [item.clause for item in clauses] == ["абзац 1", "абзац 2", "абзац 3"]
    assert clauses[2].text == "Ведёт реестр нарушений"


def test_table_only_document_has_independent_references(tmp_path: Path) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Подразделение"
    table.cell(0, 1).text = "Функция"
    table.cell(1, 0).text = "Контроль"
    table.cell(1, 1).text = "Проверяет отчёты"
    clauses = parse_docx(save_document(tmp_path, document))
    assert len(clauses) == 2
    assert clauses[1].clause == "таблица 1, строка 2"
    assert "Проверяет отчёты" in clauses[1].text


def test_body_order_and_duplicate_number_locators(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("1.1. Первая функция")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Функция из таблицы"
    document.add_paragraph("1.1. Функция приложения")
    clauses = parse_docx(save_document(tmp_path, document))
    assert [item.clause for item in clauses] == [
        "1.1 [фрагмент 1]",
        "таблица 1, строка 1",
        "1.1 [фрагмент 2]",
    ]


def test_consecutive_clauses_inside_one_paragraph() -> None:
    assert split_inline_clauses("3.9. Первая обязанность. 3.10. Вторая обязанность.") == [
        "3.9. Первая обязанность.",
        "3.10. Вторая обязанность.",
    ]


@pytest.mark.parametrize(
    "text",
    [
        "3.9. Согласно п. 3.10. действует контроль.",
        "3.9. Срок исполнения 23.09.2026.",
        "3.9. См. пункт 3.10. и раздел 4.",
    ],
)
def test_inline_references_are_not_new_clauses(text: str) -> None:
    assert split_inline_clauses(text) == [text]


def test_top_level_numbers_and_unnumbered_text(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph("1. Обязанности")
    document.add_paragraph("Функция без нумерации")
    document.add_paragraph("2) Полномочия")
    clauses = parse_docx(save_document(tmp_path, document))
    assert [item.clause for item in clauses] == ["1", "абзац 2", "2"]


def test_empty_document_is_not_silently_skipped(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="Не извлечён текст"):
        parse_documents([save_document(tmp_path, Document())])


def test_display_name_strips_only_upload_uuid() -> None:
    assert display_name(Path("a" * 32 + "_test.docx")) == "test.docx"
    name = "z" * 32 + "_test.docx"
    assert display_name(Path(name)) == name


def test_duplicate_upload_names_are_rejected(tmp_path: Path) -> None:
    paths = [tmp_path / (prefix * 32 + "_same.docx") for prefix in ("a", "b")]
    with pytest.raises(ValueError, match="Одинаковые имена"):
        parse_documents(paths)
